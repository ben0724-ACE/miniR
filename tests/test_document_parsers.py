import base64
import importlib.util
import os
import sys
import tempfile
import types
import unittest


PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
) + (b"\0" * 600)


def has_module(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


if not has_module("faiss"):
    sys.modules["faiss"] = types.ModuleType("faiss")

if not has_module("tqdm"):
    tqdm_module = types.ModuleType("tqdm")
    tqdm_module.tqdm = lambda iterable=None, *args, **kwargs: iterable if iterable is not None else []
    sys.modules["tqdm"] = tqdm_module

if not has_module("jieba"):
    jieba_module = types.ModuleType("jieba")
    jieba_module.lcut = lambda text: str(text).split()
    sys.modules["jieba"] = jieba_module

if not has_module("rank_bm25"):
    rank_bm25_module = types.ModuleType("rank_bm25")

    class BM25Okapi:
        def __init__(self, corpus):
            self.corpus = corpus

        def get_scores(self, tokens):
            return [0.0 for _ in self.corpus]

    rank_bm25_module.BM25Okapi = BM25Okapi
    sys.modules["rank_bm25"] = rank_bm25_module

if not has_module("FlagEmbedding"):
    flag_embedding = types.ModuleType("FlagEmbedding")

    class BGEM3FlagModel:
        pass

    flag_embedding.BGEM3FlagModel = BGEM3FlagModel
    sys.modules["FlagEmbedding"] = flag_embedding


from scripts.add_documents import DocumentProcessor
from scripts.db_manager import Corpus


class DocumentParserTests(unittest.TestCase):
    def make_processor(self, tmpdir: str) -> DocumentProcessor:
        return DocumentProcessor(
            md_directory=tmpdir,
            faiss_index_path=os.path.join(tmpdir, "faiss_index"),
            db_path=os.path.join(tmpdir, "rag_data.db"),
            append_mode=False,
        )

    def test_markdown_with_image(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            img_path = os.path.join(tmpdir, "pic.png")
            with open(img_path, "wb") as f:
                f.write(PNG_BYTES)
            doc_path = os.path.join(tmpdir, "note.md")
            with open(doc_path, "w", encoding="utf-8") as f:
                f.write("# Intro\nhello\n![pic](pic.png)\n")

            processor = self.make_processor(tmpdir)
            try:
                content, doc_type, images, image_map = processor.read_file(doc_path)
                sections = processor.parse_sections(content, doc_type, doc_path, images, image_map)
            finally:
                processor.db.close()

            self.assertEqual(doc_type, "markdown")
            self.assertEqual(len(images), 1)
            self.assertEqual(len(sections), 1)
            self.assertEqual(len(sections[0]["images"]), 1)

    def test_find_files_skips_generated_images_directory_and_standalone_images(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = os.path.join(tmpdir, "note.md")
            with open(doc_path, "w", encoding="utf-8") as f:
                f.write("# Intro\nhello\n")
            with open(os.path.join(tmpdir, "standalone.png"), "wb") as f:
                f.write(PNG_BYTES)
            images_dir = os.path.join(tmpdir, "images")
            os.makedirs(images_dir, exist_ok=True)
            with open(os.path.join(images_dir, "generated.png"), "wb") as f:
                f.write(PNG_BYTES)

            processor = self.make_processor(tmpdir)
            try:
                files = processor.find_files()
            finally:
                processor.db.close()

            self.assertEqual(files, [doc_path])

    def test_find_files_keeps_same_stem_different_formats(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            md_path = os.path.join(tmpdir, "same.md")
            pdf_path = os.path.join(tmpdir, "same.pdf")
            with open(md_path, "w", encoding="utf-8") as f:
                f.write("# Same\nmarkdown\n")
            with open(pdf_path, "wb") as f:
                f.write(b"%PDF-1.4\n%%EOF\n")

            processor = self.make_processor(tmpdir)
            try:
                files = processor.find_files()
            finally:
                processor.db.close()

            self.assertEqual(files, sorted([md_path, pdf_path]))

    def test_document_exists_uses_path_not_stem_only(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            md_path = os.path.join(tmpdir, "same.md")
            pdf_path = os.path.join(tmpdir, "same.pdf")
            with open(md_path, "w", encoding="utf-8") as f:
                f.write("# Same\nmarkdown\n")
            with open(pdf_path, "wb") as f:
                f.write(b"%PDF-1.4\n%%EOF\n")

            processor = self.make_processor(tmpdir)
            try:
                processor.db.add_corpus(Corpus(
                    name="same",
                    file_path=md_path,
                    relative_path=md_path,
                    chunk_count=1,
                ))

                self.assertTrue(processor._is_document_exists(md_path))
                self.assertFalse(processor._is_document_exists(pdf_path))
            finally:
                processor.db.close()

    def test_document_exists_allows_same_basename_in_different_dirs(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            dir_a = os.path.join(tmpdir, "project_a")
            dir_b = os.path.join(tmpdir, "project_b")
            os.makedirs(dir_a, exist_ok=True)
            os.makedirs(dir_b, exist_ok=True)
            first_path = os.path.join(dir_a, "spec.pdf")
            second_path = os.path.join(dir_b, "spec.pdf")
            with open(first_path, "wb") as f:
                f.write(b"%PDF-1.4\n%%EOF\n")
            with open(second_path, "wb") as f:
                f.write(b"%PDF-1.4\n%%EOF\n")

            processor = self.make_processor(tmpdir)
            try:
                processor.db.add_corpus(Corpus(
                    name="spec.pdf",
                    file_path=first_path,
                    relative_path=first_path,
                    chunk_count=1,
                ))

                self.assertTrue(processor._is_document_exists(first_path))
                self.assertFalse(processor._is_document_exists(second_path))
            finally:
                processor.db.close()

    @unittest.skipUnless(has_module("docx"), "python-docx is not installed")
    def test_word_with_heading_and_image(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmpdir:
            img_path = os.path.join(tmpdir, "pic.png")
            with open(img_path, "wb") as f:
                f.write(PNG_BYTES)
            doc_path = os.path.join(tmpdir, "guide.docx")
            doc = Document()
            doc.add_heading("Overview", level=1)
            doc.add_paragraph("Word body")
            doc.add_picture(img_path)
            doc.save(doc_path)

            processor = self.make_processor(tmpdir)
            try:
                content, doc_type, images, image_map = processor.read_file(doc_path)
                sections = processor.parse_sections(content, doc_type, doc_path, images, image_map)
            finally:
                processor.db.close()

            self.assertEqual(doc_type, "word")
            self.assertIn("Word body", content)
            self.assertGreaterEqual(len(images), 1)
            self.assertEqual(len(sections), 1)
            self.assertGreaterEqual(len(sections[0]["images"]), 1)

    @unittest.skipUnless(has_module("pptx"), "python-pptx is not installed")
    def test_pptx_multi_slide(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as tmpdir:
            img_path = os.path.join(tmpdir, "pic.png")
            with open(img_path, "wb") as f:
                f.write(PNG_BYTES)
            ppt_path = os.path.join(tmpdir, "deck.pptx")
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Slide One"
            slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1)).text = "First body"
            slide.shapes.add_picture(img_path, Inches(1), Inches(2.5), Inches(1), Inches(1))
            slide2 = prs.slides.add_slide(prs.slide_layouts[5])
            slide2.shapes.title.text = "Slide Two"
            slide2.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1)).text = "Second body"
            prs.save(ppt_path)

            processor = self.make_processor(tmpdir)
            try:
                content, doc_type, images, image_map = processor.read_file(ppt_path)
                sections = processor.parse_sections(content, doc_type, ppt_path, images, image_map)
            finally:
                processor.db.close()

            self.assertEqual(doc_type, "ppt")
            self.assertIn("First body", content)
            self.assertIn("Second body", content)
            self.assertEqual(len(sections), 2)
            self.assertGreaterEqual(len(images), 1)

    @unittest.skipUnless(has_module("openpyxl"), "openpyxl is not installed")
    def test_excel_multi_sheet(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = os.path.join(tmpdir, "book.xlsx")
            workbook = Workbook()
            sheet1 = workbook.active
            sheet1.title = "Alpha"
            sheet1["A1"] = "Name"
            sheet1["B1"] = "Value"
            sheet1["A2"] = "foo"
            sheet1["B2"] = 42
            sheet2 = workbook.create_sheet("Beta")
            sheet2["A1"] = "Other"
            sheet2["A2"] = "bar"
            workbook.save(xlsx_path)

            processor = self.make_processor(tmpdir)
            try:
                content, doc_type, images, image_map = processor.read_file(xlsx_path)
                sections = processor.parse_sections(content, doc_type, xlsx_path, images, image_map)
            finally:
                processor.db.close()

            self.assertEqual(doc_type, "excel")
            self.assertIn("工作表：Alpha", content)
            self.assertIn("工作表：Beta", content)
            self.assertEqual(len(sections), 2)

    @unittest.skipUnless(has_module("fitz"), "PyMuPDF is not installed")
    def test_pdf_text_page(self):
        import fitz

        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = os.path.join(tmpdir, "file.pdf")
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "PDF body")
            pdf.save(pdf_path)
            pdf.close()

            processor = self.make_processor(tmpdir)
            try:
                content, doc_type, images, image_map = processor.read_file(pdf_path)
                sections = processor.parse_sections(content, doc_type, pdf_path, images, image_map)
            finally:
                processor.db.close()

            self.assertEqual(doc_type, "pdf")
            self.assertIn("PDF body", content)
            self.assertEqual(len(sections), 1)

if __name__ == "__main__":
    unittest.main()
